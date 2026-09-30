"""Generate training questions (v4) for every Legal_Knowledge_Base_v2.xlsx section that has none yet.

Same rules as scripts/generate_training_pairs_v3.py (plain-language questions from a citizen's point of view; no Act names,
no section numbers, no legal jargon; >=2 short casual questions; <=25 words; retry with feedback; definitions auto-skipped),
plus: model openai/gpt-oss-120b (reasoning_effort low, temperature 0.7), parallel workers under a token-rate limit, a persistent cost
ledger with a hard stop, boilerplate sections auto-skipped, some questions with typos / Hinglish-style English, and a mechanical
check that a question does not copy 5+ consecutive words of the section.

    python scripts/generate_training_pairs_v4.py --plan              # size the job, no API calls
    python scripts/generate_training_pairs_v4.py --limit 20          # small trial
    python scripts/generate_training_pairs_v4.py --workers 6         # full run (resumable: just run it again)

Outputs (all resumable): data/training_pairs_v4.jsonl, data/training_pairs_v4_skipped.csv, data/training_pairs_v4_cost.json
"""
import argparse
import ast
import collections
import csv
import json
import os
import re
import sys
import threading
import time
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime

import openpyxl

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(BASE_DIR, ".."))
sys.path.insert(0, BASE_DIR)
DATA = os.path.join(ROOT, "data")
KB_EXCEL = os.path.join(ROOT, "Legal_Knowledge_Base_v2.xlsx")
OUTPUT_JSONL = os.path.join(DATA, "training_pairs_v4.jsonl")
SKIPPED_CSV = os.path.join(DATA, "training_pairs_v4_skipped.csv")
COST_JSON = os.path.join(DATA, "training_pairs_v4_cost.json")
# sections that already have questions (any attempt) or were deliberately skipped earlier
EXISTING_JSONL = ["training_pairs_old_clean.jsonl", "training_pairs_v2.jsonl", "training_pairs_v2_clean.jsonl", "training_pairs_v3.jsonl"]
EXISTING_SKIPPED = ["training_pairs_v2_skipped.csv", "training_pairs_v3_skipped.csv"]

MODEL_NAME = "openai/gpt-oss-120b"
TEMPERATURE = 0.7
MIN_SHORT_QUESTIONS = 2   # at least this many questions must be short and casual (as in v3)
SHORT_MAX_WORDS = 10
COPY_NGRAM = 5            # a question may not repeat this many consecutive words of the section
# Groq list price for openai/gpt-oss-120b (USD per million tokens). The cost shown is an estimate from the usage the API reports;
# output is priced at the higher of the published figures so the stop-at-budget guard errs on the safe side.
PRICE_IN_PER_M, PRICE_OUT_PER_M = 0.15, 0.75
MAX_COST_USD = 10.0
TPM_BUDGET = 200_000      # stay under the 250K tokens/minute limit


# ---------------------------------------------------------------- v3 rules (unchanged unless noted)
DEFINITION_TITLE = re.compile(r"^\s*[“‘\"']|\bdefinitions?\b|\binterpretation\b|\bmeaning of\b", re.IGNORECASE)
DEFINITION_TEXT = re.compile(r"in this (?:act|code|adhiniyam),?\s*unless the context otherwise requires", re.IGNORECASE)


def is_definition_section(title, text):
    if DEFINITION_TITLE.search(title or ""):
        return True
    return DEFINITION_TEXT.search((text or "")[:300]) is not None


def build_system_prompt(n):
    return f"""You generate questions that an ordinary Indian citizen with NO legal knowledge might type into a search box when facing a real-life problem that this specific legal section would answer.

First, understand the section: in ONE plain sentence, say what this section lets a citizen do, protects them from, or what happens to them. Then write questions about exactly that situation.

Rules for the questions:
- Generate exactly {n} different questions describing different real-life situations.
- At least {MIN_SHORT_QUESTIONS} questions must be SHORT and CASUAL: {SHORT_MAX_WORDS} words or fewer, like a quick phone search (e.g. "cheque bounced what can i do", "boss not paying salary").
- Make a couple of the questions sound like real typing: lowercase, a small typo, or Hinglish-style English (e.g. "landlord deposit wapas nahi de raha"). The rest should be normal sentences.
- The other questions can describe the situation in one or two sentences, under 25 words.
- Use only everyday words. Do NOT use legal or technical terms from the section (for example: holder in due course, drawee, instrument, endorsement, co-parcener, cognizance, presumption). Describe the situation instead.
- Do NOT copy phrases from the section text; put the situation in your own words.
- Talk only about your own situation, in the first person.
- NEVER mention the name of any law, 'this Act', 'the Act', 'this Code', any abbreviation of a law, or any section number.
- Every question must be one that THIS section actually answers, according to your one-sentence summary.
- SKIP the section if it is purely technical (short title, extent, commencement, repeal, power to make rules, internal court procedure) and no citizen would ask about it. To skip, return: {{"summary": "", "questions": [], "skip_reason": "short explanation"}}
- Otherwise return ONLY valid JSON: {{"summary": "one plain sentence", "questions": ["question 1", "question 2", ...]}}
"""


BANNED_PHRASES = [
    "this act", "the act", "this law", "this code", "sanhita", "adhiniyam",
    "domestic violence act", "hindu marriage act", "special marriage act",
    "hindu succession act", "negotiable instruments", "code on wages",
    "code on social security", "dowry prohibition act", "senior citizens act",
    "holder in due course",
]
BANNED_WORDS = ["pocso", "posh", "bsa", "bns", "bnss"]


def normalize_sec(sec):
    if sec is None:
        return ""
    s = str(sec).strip()
    return s[:-2] if s.endswith(".0") else s


def words(s):
    return re.findall(r"[a-z0-9']+", (s or "").lower())


def ngrams(ws, n):
    return {tuple(ws[i:i + n]) for i in range(len(ws) - n + 1)}


def is_valid_question(q, section_ngrams=None):
    if not q or not isinstance(q, str):
        return False, "Empty or invalid string"
    if re.search(r"\bAct\b", q):
        return False, "Contains 'Act'"
    q_lower = q.lower()
    for phrase in BANNED_PHRASES:
        if phrase in q_lower:
            return False, f"Contains '{phrase}'"
    for word in BANNED_WORDS:
        if re.search(rf"\b{word}\b", q_lower):
            return False, f"Contains '{word}'"
    if re.search(r"\bsection\s*\d+\b", q, re.IGNORECASE):
        return False, "Contains section number"
    if re.search(r"\bsec\.\s*\d+\b", q, re.IGNORECASE):
        return False, "Contains sec. number"
    if len(q.strip().split()) > 25:
        return False, "Over 25 words"
    if section_ngrams and ngrams(words(q), COPY_NGRAM) & section_ngrams:   # added in v4: no copied legal wording
        return False, f"Copies {COPY_NGRAM}+ consecutive words from the section text"
    return True, "OK"


def count_short(questions):
    return sum(1 for q in questions if len(q.strip().split()) <= SHORT_MAX_WORDS)


# ---------------------------------------------------------------- app placeholder filter, without importing torch
def load_placeholder_filter():
    src = open(os.path.join(BASE_DIR, "search_core.py"), encoding="utf-8").read()
    node = next(n for n in ast.parse(src).body if isinstance(n, ast.FunctionDef) and n.name == "is_placeholder_record")
    ns = {}
    exec(compile(ast.Module(body=[node], type_ignores=[]), "search_core.is_placeholder_record", "exec"), ns)
    return ns["is_placeholder_record"]


# ---------------------------------------------------------------- cost ledger and rate limiter
class Ledger:
    """Cumulative usage across runs (persisted), plus a sliding one-minute token window for the rate limit."""

    def __init__(self):
        self.lock = threading.Lock()
        self.window = collections.deque()
        d = json.load(open(COST_JSON, encoding="utf-8")) if os.path.exists(COST_JSON) else {}
        self.prompt_tokens = d.get("prompt_tokens", 0)
        self.completion_tokens = d.get("completion_tokens", 0)
        self.calls = d.get("calls", 0)
        self.stopped = threading.Event()

    @property
    def cost(self):
        return self.prompt_tokens * PRICE_IN_PER_M / 1e6 + self.completion_tokens * PRICE_OUT_PER_M / 1e6

    def wait_for_room(self, est_tokens=1600):
        while True:
            with self.lock:
                now = time.time()
                while self.window and now - self.window[0][0] > 60:
                    self.window.popleft()
                used = sum(t for _, t in self.window)
                if used + est_tokens <= TPM_BUDGET or self.stopped.is_set():
                    return
                wait = 60 - (now - self.window[0][0]) + 0.2
            time.sleep(max(0.2, min(wait, 5)))

    def add(self, usage):
        pt = int(getattr(usage, "prompt_tokens", 0) or 0)
        ct = int(getattr(usage, "completion_tokens", 0) or 0)
        with self.lock:
            self.prompt_tokens += pt
            self.completion_tokens += ct
            self.calls += 1
            self.window.append((time.time(), pt + ct))
            if self.cost >= MAX_COST_USD:
                self.stopped.set()

    def save(self):
        with self.lock:
            json.dump({"prompt_tokens": self.prompt_tokens, "completion_tokens": self.completion_tokens, "calls": self.calls,
                       "estimated_cost_usd": round(self.cost, 4), "model": MODEL_NAME,
                       "price_per_million_usd": {"input": PRICE_IN_PER_M, "output": PRICE_OUT_PER_M},
                       "updated_at": datetime.now().isoformat()}, open(COST_JSON, "w", encoding="utf-8"), indent=2)


def call_groq_api(client, ledger, n, act_name, section_number, section_title, legal_text):
    text_snippet = legal_text[:1000] if legal_text else ""
    section_ngrams = ngrams(words(section_title + " " + legal_text), COPY_NGRAM)
    base_user_prompt = (f"Act: {act_name}\nSection Number: {section_number}\n"
                        f"Section Title: {section_title}\nLegal Text: {text_snippet}")
    system_prompt = build_system_prompt(n)
    backoff_delays = [10, 20, 40, 60, 60, 60, 60, 60]
    feedback, parsed_result = "", None
    for gen_pass in range(3):
        if ledger.stopped.is_set():
            return None, "BUDGET_STOP"
        user_prompt = base_user_prompt
        if feedback:
            user_prompt += ("\n\nCRITICAL FIX NEEDED in your previous answer:\n"
                            f"{feedback}\nPlease regenerate {n} questions following ALL rules strictly.")
        attempt_success = False
        for attempt in range(8):
            ledger.wait_for_room()
            if ledger.stopped.is_set():
                return None, "BUDGET_STOP"
            try:
                response = client.chat.completions.create(
                    model=MODEL_NAME,
                    messages=[{"role": "system", "content": system_prompt}, {"role": "user", "content": user_prompt}],
                    temperature=TEMPERATURE, max_tokens=1500, reasoning_effort="low")
                ledger.add(response.usage)
                raw = (response.choices[0].message.content or "").strip()
                if raw.startswith("```"):
                    raw = raw.split("```")[1]
                    raw = raw[4:] if raw.startswith("json") else raw
                    raw = raw.strip()
                parsed_result = json.loads(raw)
                attempt_success = True
                break
            except Exception as e:
                err = str(e)
                low = err.lower()
                if "tokens per day" in low or "tpd" in low:
                    return None, "DAILY_LIMIT"
                if "429" in err or "rate limit" in low or "tpm" in low or "rpm" in low:
                    m = re.search(r"try again in (\d+(?:\.\d+)?)s", err, re.IGNORECASE)
                    time.sleep(float(m.group(1)) + 1.0 if m else backoff_delays[attempt])
                    continue
                if attempt < 7:
                    time.sleep(2)
                    continue
                return None, f"ERROR: {err[:200]}"
        if not attempt_success or parsed_result is None:
            return None, "MAX_RETRIES_EXCEEDED"

        questions = parsed_result.get("questions", [])
        skip_reason = parsed_result.get("skip_reason", "")
        if not questions or skip_reason:
            return parsed_result, None
        invalid_feedback, valid_questions = [], []
        for q in questions:
            ok, reason = is_valid_question(q, section_ngrams)
            if ok:
                valid_questions.append(q)
            else:
                invalid_feedback.append(f"- Question '{q}': {reason}")
        if count_short(valid_questions) < MIN_SHORT_QUESTIONS:
            invalid_feedback.append(f"- Only {count_short(valid_questions)} question(s) are short and casual; at least "
                                    f"{MIN_SHORT_QUESTIONS} must be {SHORT_MAX_WORDS} words or fewer, in everyday words.")
        if not invalid_feedback:
            return {"summary": parsed_result.get("summary", ""), "questions": valid_questions, "skip_reason": ""}, None
        feedback = "\n".join(invalid_feedback)
    final_questions = [q for q in parsed_result.get("questions", []) if is_valid_question(q, section_ngrams)[0]]
    return {"summary": parsed_result.get("summary", ""), "questions": final_questions,
            "skip_reason": parsed_result.get("skip_reason", "")}, None


# ---------------------------------------------------------------- planning
def load_kb_sections():
    """One entry per (act, section): the KB v2 row (first part of a split section), minus rows the app itself would filter out."""
    is_placeholder = load_placeholder_filter()
    wb = openpyxl.load_workbook(KB_EXCEL, read_only=True, data_only=True)
    ws = wb.active
    it = ws.iter_rows(values_only=True)
    headers = [str(h).strip() if h is not None else "" for h in next(it)]
    ix = {h: i for i, h in enumerate(headers)}
    sections, placeholders = collections.OrderedDict(), 0
    for row in it:
        rec = {h: row[i] for h, i in ix.items()}
        bad, _ = is_placeholder(rec)
        if bad:
            placeholders += 1
            continue
        act, sec = str(rec["act_name"]).strip(), normalize_sec(rec["section_number"])
        title = re.sub(r"\s*\(part \d+ of \d+\)$", "", str(rec.get("section_title") or ""))
        key = (act, sec)
        if key not in sections:                       # a section split into parts is generated once, from its first part
            sections[key] = {"act_name": act, "section_number": sec, "section_title": title,
                             "legal_text": str(rec.get("legal_text") or ""), "source": rec.get("source")}
    wb.close()
    return sections, placeholders


def load_existing():
    done, skipped = set(), set()
    for fn in EXISTING_JSONL + [os.path.basename(OUTPUT_JSONL)]:
        p = os.path.join(DATA, fn)
        if os.path.exists(p):
            for line in open(p, encoding="utf-8"):
                if line.strip():
                    try:
                        d = json.loads(line)
                        done.add((str(d.get("act_name", "")).strip(), normalize_sec(d.get("section_number"))))
                    except Exception:
                        pass
    for fn in EXISTING_SKIPPED + [os.path.basename(SKIPPED_CSV)]:
        p = os.path.join(DATA, fn)
        if os.path.exists(p):
            for r in csv.DictReader(open(p, encoding="utf-8")):
                skipped.add((str(r.get("act_name", "")).strip(), normalize_sec(r.get("section_number"))))
    return done, skipped


def auto_skip_reason(title, text):
    from kb_v2_variants_eval import section_type  # same section-type patterns as the KB v2 diagnosis
    if is_definition_section(title, text):
        return "Definition section (auto-skipped)"
    ty = section_type(title)
    if ty == "short title / extent / commencement":
        return "Short title / extent / commencement (auto-skipped)"
    if ty == "repeal / savings":
        return "Repeal / savings (auto-skipped)"
    if ty == "rule-making power":
        return "Rule-making power (auto-skipped)"
    return None


def plan():
    sections, placeholders = load_kb_sections()
    done, skipped = load_existing()
    todo, auto = [], collections.Counter()
    auto_rows = []
    already = prev_skipped = 0
    for key, s in sections.items():
        if key in done:
            already += 1
        elif key in skipped:
            prev_skipped += 1
        else:
            why = auto_skip_reason(s["section_title"], s["legal_text"])
            if why:
                auto[why] += 1
                auto_rows.append((s, why))
            else:
                todo.append(s)
    return sections, placeholders, already, prev_skipped, auto, auto_rows, todo


def log_skip(act, sec, title, reason):
    new = not os.path.exists(SKIPPED_CSV)
    with open(SKIPPED_CSV, "a", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        if new:
            w.writerow(["act_name", "section_number", "section_title", "skip_reason", "logged_at"])
        w.writerow([act, sec, title, reason, datetime.now().isoformat()])


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--n", type=int, default=5)
    ap.add_argument("--workers", type=int, default=6)
    ap.add_argument("--limit", type=int, default=None, help="only process this many sections (trial run)")
    ap.add_argument("--plan", action="store_true", help="print the job size and exit (no API calls)")
    ap.add_argument("--act", type=str, default=None, help="only this exact act_name")
    args = ap.parse_args()

    sections, placeholders, already, prev_skipped, auto, auto_rows, todo = plan()
    if args.act:
        todo = [s for s in todo if s["act_name"] == args.act]
    print(f"KB v2 sections (one per act+section, placeholders removed): {len(sections)}  (placeholder rows filtered: {placeholders})")
    print(f"  already have questions: {already} | skipped in earlier runs: {prev_skipped} | auto-skipped now: {sum(auto.values())} {dict(auto)}")
    print(f"  TO GENERATE: {len(todo)} sections x {args.n} questions")
    per_act = collections.Counter(s["act_name"] for s in todo)
    print("  biggest Acts:", per_act.most_common(6))
    if args.plan:
        est = len(todo) * (1000 * PRICE_IN_PER_M + 700 * PRICE_OUT_PER_M) / 1e6
        print(f"  rough cost estimate: ~${est:.2f} (about 1,000 input + 700 output tokens per section; retries add ~15%)")
        return

    # auto-skips are logged without any API call
    for s, why in auto_rows:
        log_skip(s["act_name"], s["section_number"], s["section_title"], why)
    if args.limit:
        todo = todo[:args.limit]

    from dotenv import load_dotenv
    from groq import Groq
    load_dotenv()
    api_key = os.environ.get("GROQ_API_KEY_GEN") or os.environ.get("GROQ_API_KEY")
    if not api_key:
        print("Error: no GROQ_API_KEY_GEN or GROQ_API_KEY in the environment/.env.")
        sys.exit(1)
    client = Groq(api_key=api_key)
    ledger = Ledger()
    print(f"Cost so far (earlier runs): ${ledger.cost:.4f}. Hard stop at ${MAX_COST_USD:.2f}. Workers: {args.workers}. Model: {MODEL_NAME}")

    out_lock = threading.Lock()
    counts = collections.Counter()
    t0 = time.time()

    def work(item):
        if ledger.stopped.is_set():
            return
        act, sec, title = item["act_name"], item["section_number"], item["section_title"]
        parsed, err = call_groq_api(client, ledger, args.n, act, sec, title, item["legal_text"])
        with out_lock:
            counts["attempted"] += 1
            if err in ("BUDGET_STOP", "DAILY_LIMIT"):
                counts["stopped_" + err] += 1
                if err == "DAILY_LIMIT":
                    ledger.stopped.set()
            elif parsed is None:
                counts["failed"] += 1
                print(f"  FAILED {act} {sec}: {err}", flush=True)
            else:
                questions, skip_reason = parsed.get("questions", []), parsed.get("skip_reason", "")
                if not questions or skip_reason:
                    counts["skipped_by_model"] += 1
                    log_skip(act, sec, title, str(skip_reason or "No questions generated").strip())
                else:
                    counts["sections_done"] += 1
                    counts["questions"] += len(questions)
                    now_iso = datetime.now().isoformat()
                    with open(OUTPUT_JSONL, "a", encoding="utf-8") as f:
                        for q in questions:
                            f.write(json.dumps({"query": q, "act_name": act, "section_number": sec, "section_title": title,
                                                "section_summary": parsed.get("summary", ""), "model": MODEL_NAME,
                                                "generated_at": now_iso}, ensure_ascii=False) + "\n")
            if counts["attempted"] % 25 == 0:
                ledger.save()
                el = time.time() - t0
                print(f"[{counts['attempted']}/{len(todo)}] done {counts['sections_done']} sections / {counts['questions']} questions | "
                      f"skipped {counts['skipped_by_model']} failed {counts['failed']} | tokens in {ledger.prompt_tokens:,} out {ledger.completion_tokens:,} | "
                      f"cost ${ledger.cost:.3f} | {counts['attempted'] / el * 60:.0f} sections/min", flush=True)

    with ThreadPoolExecutor(max_workers=args.workers) as ex:
        list(ex.map(work, todo))
    ledger.save()
    print(f"\nFinished this run: {dict(counts)} | cumulative cost ${ledger.cost:.4f} ({ledger.calls} calls)"
          + ("  ** STOPPED: budget reached **" if ledger.cost >= MAX_COST_USD else ""))


if __name__ == "__main__":
    main()
