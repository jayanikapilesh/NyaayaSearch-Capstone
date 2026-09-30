import os
import re
import shutil
from openpyxl import load_workbook

KB = "Legal_Knowledge_Base_v2.xlsx"
BACKUP = "Legal_Knowledge_Base_v2_before_ica_fix.xlsx"
TEXT = "data/processed/indian_contract_act_1872.txt"
ACT = "Indian Contract Act, 1872"
SOURCE_DATE_CHECKED = "2026-09-30"

# --- Same parsing rules as scripts/parse_new_acts.py ---
FOOTNOTE_OPENINGS = (
    r"The\s+words?|The\s+Act\s+has\s+been\s+extended|Certain\s+words|Subs\s*\.|Ins\s*\.|ins\s*\.|"
    r"Omitted|omitted|Inserted|Substituted|subs\s*\.|Cl\s*\.|Section\s+\d+\s+substituted|"
    r"Section\s+\d+\s+numbered|vide\s+notification|w\s*\.\s*e\s*\.\s*f|Rep\s*\.|rep\s*\.|Added|added"
)
body_anchor = re.compile(r"\.\s*[\u2014\u2013\u2015\-]+\s*(?:\(\s*1\s*\)\s*)?This\s+(?:Act|Code)\s+may\s+be\s+call\w*\s*d", re.DOTALL)
section_start_pattern = re.compile(r"\n(\d{1,3}[A-Z]?)\.\s")
footnote_pattern = re.compile(r"\n\s*\d{1,2}\.\s+(?:" + FOOTNOTE_OPENINGS + r").*?(?=\n\d|\n\n|\Z)", re.DOTALL)
footnote_like = re.compile(r"^\d?\[?\d{1,3}[A-Z]?\.\s*(?:" + FOOTNOTE_OPENINGS + r")")
section_pattern = re.compile(r"(?<!\d)(\d{1,3}[A-Z]?)\.\s*(?:\u2014|\u2013|\u2015|\-)?\s*(?=[A-Z\u201c\[])|(?<!\d)\d\[(\d{1,3}[A-Z]?)\.\s+")
xref_before = re.compile(r"(?:sections?|sub-sections?|clauses?|articles?|rules?|\band|\bor|,)\s*$", re.IGNORECASE)


def is_junk(t):
    return len(t) < 60 or footnote_like.match(t) is not None


def parse_sections(text):
    anchor = body_anchor.search(text)
    starts = list(section_start_pattern.finditer(text[:anchor.start()]))
    text2 = text[starts[-1].start() + 1:]
    end = len(text2)
    for marker in ["SCHEDULE", "THE SCHEDULE", "THE FIRST SCHEDULE"]:
        pos = text2.find(marker)
        if pos != -1:
            end = min(end, pos)
    text2 = footnote_pattern.sub("\n", text2[:end])
    matches = [m for m in section_pattern.finditer(text2)
               if not xref_before.search(text2[max(0, m.start() - 30):m.start()])]
    sections = {}
    for i, m in enumerate(matches):
        label = m.group(1) or m.group(2)
        s_end = matches[i + 1].start() if i + 1 < len(matches) else len(text2)
        t = text2[m.start():s_end].strip()
        t = re.sub(r"\n\s*\d+\s*\n", "\n", t).replace("IndiaCode", " ")
        t = re.sub(r"\s+", " ", t).strip()
        if label not in sections or (is_junk(sections[label]) and not is_junk(t)):
            sections[label] = t
    return sections


def split_title(label, full_text):
    content = re.sub(rf"^\d?\[?{re.escape(label)}\.\s*(?:\u2014|\u2013|\u2015|\-)?\s*", "", full_text, count=1).strip()
    m = re.match(r"^(.+?)(?:\u2014|\u2013|\u2015)\s*", content)
    if not m:
        m = re.match(r"^(.+?\.)\s", content)
    title = m.group(1).strip() if m else content[:150]
    return title, content


# --- 1. Backup once; always rebuild from the backup (safe to re-run) ---
if not os.path.exists(BACKUP):
    shutil.copyfile(KB, BACKUP)
    print("Backup created:", BACKUP)
wb = load_workbook(BACKUP)
ws = wb.active
headers = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
col = {h: i for i, h in enumerate(headers)}

have = set()
for row in ws.iter_rows(min_row=2, values_only=True):
    if str(row[col["act_name"]]).strip() == ACT:
        have.add(str(row[col["section_number"]]).strip())

# --- 2. Re-parse the Contract Act and pick the missing, usable sections ---
with open(TEXT, "r", encoding="utf-8") as f:
    text = f.read()

# Footnote marker glued to the section number in the PDF ("1151." -> "151.", "1161." -> "161.")
text = re.sub(r"\n\s*\d(151\.\s*Care to be taken)", r"\n\1", text)
text = re.sub(r"\n\s*\d(161\.\s*Bailee)", r"\n\1", text)
parsed = parse_sections(text)

added = []
for label in sorted(parsed, key=lambda x: (int(re.match(r"\d+", x).group()), x)):
    num = int(re.match(r"\d+", label).group())
    if label in have or 76 <= num <= 123:      # already present, or repealed (moved to Sale of Goods Act, 1930)
        continue
    title, content = split_title(label, parsed[label])
    if len(content) < 100 or content.strip() == title.strip():
        continue                                # skip broken/empty text
    values = {
        "law_id": f"ICA-{label}", "domain": "Contract", "act_name": ACT, "act_number": "9 of 1872",
        "section_number": label, "section_title": title, "legal_text": content,
        "jurisdiction": "India", "enforcement_status": "To be verified", "effective_date": "",
        "source_url": "https://www.indiacode.nic.in/", "source_date_checked": SOURCE_DATE_CHECKED,
        "verification_status": "Candidate\u2014verify",
    }
    ws.append([values.get(h, "") for h in headers])
    added.append((label, title[:50], len(content)))

wb.save(KB)
print(f"Contract Act sections before: {len(have)} | added: {len(added)} | now: {len(have) + len(added)}")
for s in ["16", "43", "70", "71", "73", "151", "160", "161"]:
    hit = [a for a in added if a[0] == s]
    print(f"  Sec {s}: " + (f"ADDED | {hit[0][1]} | {hit[0][2]} chars" if hit else ("already there" if s in have else "NOT FOUND")))
print("Saved to:", KB)