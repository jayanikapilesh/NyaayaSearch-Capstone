import os
import shutil
from collections import Counter
from openpyxl import load_workbook, Workbook

COMBINED = "Legal_Knowledge_Base_combined.xlsx"
BASE = "Legal_Knowledge_Base_combined_before_new_acts.xlsx"  # backup of the 11-Act version

NEW_FILES = [
    "Legal_Knowledge_Base_hma.xlsx",
    "Legal_Knowledge_Base_hsa.xlsx",
    "Legal_Knowledge_Base_sma.xlsx",
    "Legal_Knowledge_Base_dpa.xlsx",
    "Legal_Knowledge_Base_hama.xlsx",
    "Legal_Knowledge_Base_nia.xlsx",
    "Legal_Knowledge_Base_cow.xlsx",
    "Legal_Knowledge_Base_css.xlsx",
    "Legal_Knowledge_Base_posh.xlsx",
    "Legal_Knowledge_Base_mwpsc.xlsx",
    "Legal_Knowledge_Base_pocso.xlsx",
    "Legal_Knowledge_Base_bsa.xlsx",
]


def read_rows(path):
    wb = load_workbook(path, read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    return list(rows[0]), rows[1:]


# 1. Back up the current knowledge base once, then always build from that backup.
if not os.path.exists(BASE):
    shutil.copyfile(COMBINED, BASE)
    print("Backup created:", BASE)
else:
    print("Using existing backup:", BASE)

headers, base_rows = read_rows(BASE)
base_acts = {r[2] for r in base_rows}

out = Workbook()
sheet = out.active
sheet.title = "Legal Knowledge Base"
sheet.append(headers)
for r in base_rows:
    sheet.append(list(r))

all_ids = [r[0] for r in base_rows]
new_total = 0

# 2. Add each new Act, with safety checks.
for f in NEW_FILES:
    h, rows = read_rows(f)
    if h != headers:
        raise SystemExit(f"STOP: column headers in {f} don't match the knowledge base.")
    act = rows[0][2]
    if act in base_acts:
        raise SystemExit(f"STOP: {act} is already in the knowledge base.")
    for r in rows:
        sheet.append(list(r))
        all_ids.append(r[0])
    new_total += len(rows)
    print(f"{f}: {len(rows)} rows ({act})")

# 3. Report and save.
dups = [k for k, v in Counter(all_ids).items() if v > 1 and k]
print()
print("Existing rows:", len(base_rows), "| New rows:", new_total, "| Total:", len(base_rows) + new_total)
print("Duplicate law_ids:", len(dups), dups[:10])

out.save(COMBINED)
print("Saved to:", COMBINED)