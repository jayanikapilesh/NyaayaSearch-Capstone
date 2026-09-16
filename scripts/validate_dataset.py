from openpyxl import load_workbook

file_path = "Legal_Knowledge_Base_v1.xlsx"

workbook = load_workbook(file_path)
sheet = workbook.active

print("🔍 Starting dataset validation...\n")

# Read column names
headers = [cell.value for cell in sheet[1]]

print("Columns found:", len(headers))

# Required columns
required_columns = [
    "law_id",
    "domain",
    "act_name",
    "section_number",
    "section_title",
    "legal_text",
    "jurisdiction",
    "source_url",
    "verification_status"
]

print("\nChecking required columns...")

for column in required_columns:
    if column in headers:
        print(f"✅ {column}")
    else:
        print(f"❌ Missing: {column}")

# Check for duplicate law IDs
law_id_column = headers.index("law_id") + 1

law_ids = []
duplicates = []

for row in range(2, sheet.max_row + 1):
    law_id = sheet.cell(row=row, column=law_id_column).value

    if law_id in law_ids:
        duplicates.append(law_id)

    law_ids.append(law_id)

print("\nChecking duplicate IDs...")

if duplicates:
    print("❌ Duplicate IDs:", duplicates)
else:
    print("✅ No duplicate law IDs")

# Check empty legal text
legal_text_column = headers.index("legal_text") + 1

empty_legal_text = []

for row in range(2, sheet.max_row + 1):
    value = sheet.cell(row=row, column=legal_text_column).value

    if not value:
        empty_legal_text.append(row)

print("\nChecking legal text...")

print("Rows without legal text:", len(empty_legal_text))

print("\n🎉 Validation complete!")