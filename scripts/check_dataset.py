from openpyxl import load_workbook

file_path = "Legal_Knowledge_Base_v1.xlsx"

workbook = load_workbook(file_path)
sheet = workbook.active

print("✅ Excel file opened successfully!")
print("Sheet name:", sheet.title)
print("Rows:", sheet.max_row)
print("Columns:", sheet.max_column)

print("\n📋 Column names:")

for cell in sheet[1]:
    print("-", cell.value)