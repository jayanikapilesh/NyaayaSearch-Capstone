from openpyxl import load_workbook

file_path = "Legal_Knowledge_Base_107.xlsx"

workbook = load_workbook(file_path)
sheet = workbook.active

print("Starting legal knowledge base validation...\n")

# 1. Check row count
data_rows = sheet.max_row - 1
print("Sections:", data_rows)

if data_rows == 107:
    print("✅ Exactly 107 sections found")
else:
    print("❌ Expected 107 sections")

# 2. Check required columns
headers = [cell.value for cell in sheet[1]]

required_columns = [
    "law_id",
    "domain",
    "act_name",
    "act_number",
    "section_number",
    "section_title",
    "legal_text",
    "jurisdiction",
    "source_url",
    "verification_status"
]

print("\nChecking required columns...")

for column in required_columns:
    if column in headers:
        print(f"✅ {column}")
    else:
        print(f"❌ Missing: {column}")

# 3. Check section numbers
section_column = headers.index("section_number") + 1
section_numbers = []

for row in range(2, sheet.max_row + 1):
    section_numbers.append(
        sheet.cell(row=row, column=section_column).value
    )

expected_sections = list(range(1, 108))

if section_numbers == expected_sections:
    print("\n✅ Section numbers 1–107 are complete and ordered")
else:
    print("\n❌ Section numbering problem")

# 4. Check duplicate law IDs
law_id_column = headers.index("law_id") + 1
law_ids = [
    sheet.cell(row=row, column=law_id_column).value
    for row in range(2, sheet.max_row + 1)
]

if len(law_ids) == len(set(law_ids)):
    print("✅ No duplicate law IDs")
else:
    print("❌ Duplicate law IDs found")

# 5. Check legal text
legal_text_column = headers.index("legal_text") + 1
missing_text = []

for row in range(2, sheet.max_row + 1):
    value = sheet.cell(row=row, column=legal_text_column).value

    if not value or len(value.strip()) < 50:
        missing_text.append(
            sheet.cell(row=row, column=section_column).value
        )

if not missing_text:
    print("✅ All sections contain legal text")
else:
    print("❌ Missing/short legal text:", missing_text)

print("\nValidation complete!")