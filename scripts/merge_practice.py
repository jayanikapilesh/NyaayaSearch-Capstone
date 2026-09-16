from openpyxl import load_workbook, Workbook

files = [
    "Legal_Knowledge_Base_verified.xlsx",
    "Legal_Knowledge_Base_property.xlsx",
    "Legal_Knowledge_Base_bns.xlsx",
    "Legal_Knowledge_Base_bnss.xlsx",
    "Legal_Knowledge_Base_contract.xlsx",
    "Legal_Knowledge_Base_dv.xlsx",
    "Legal_Knowledge_Base_specific_relief.xlsx",
    "Legal_Knowledge_Base_it.xlsx"
]

output_file = "Legal_Knowledge_Base_combined.xlsx"

workbook = Workbook()
sheet = workbook.active
sheet.title = "Legal Knowledge Base"

headers = [
    "law_id",
    "domain",
    "act_name",
    "act_number",
    "section_number",
    "section_title",
    "legal_text",
    "jurisdiction",
    "enforcement_status",
    "effective_date",
    "source_url",
    "source_date_checked",
    "verification_status"
]

sheet.append(headers)

total = 0

for file in files:
    wb = load_workbook(file, read_only=True)
    ws = wb.active
    count = 0

    for row in ws.iter_rows(min_row=2, values_only=True):
        sheet.append(row)
        count += 1
        total += 1

    wb.close()
    print(f"{file}: {count} rows")

workbook.save(output_file)

print("Total rows:", total)
print("Saved to:", output_file)