"""Build Legal_Knowledge_Base_v2.xlsx: our own 23 Acts (copied verbatim) + cleaned Acts from Open India Law.

Offline dataset-construction script (not part of the request path). Run from the repo root:
    python scripts/build_kb_v2.py

Inputs : Legal_Knowledge_Base_combined.xlsx           (never modified)
         data/open_india_law/act_quality_report.csv   (from act_quality_check.py)
         data/open_india_law/sections_fixed.parquet   (folded sections re-split, one row per section)
         data/open_india_law/in_*_legislation.parquet (only for act number / source url)
Outputs: Legal_Knowledge_Base_v2.xlsx
         data/open_india_law/kb_v2_build_report.json, kb_v2_acts.csv, kb_v2_dropped_sections.csv
"""
import collections
import datetime
import html
import json
import re
from pathlib import Path

import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parent.parent
OIL = ROOT / "data" / "open_india_law"
OLD_KB = ROOT / "Legal_Knowledge_Base_combined.xlsx"
NEW_KB = ROOT / "Legal_Knowledge_Base_v2.xlsx"
SOURCE_OIL = "open-india-law v2026.08, CC BY 4.0"
TODAY = datetime.date.today().isoformat()

# Indian Stamp Act and Industrial Disputes Act were deliberately removed from this list: a full PDF check verified
# only 5% / 36% of their rows (text sits under the wrong section numbers).
# NEEDS CHECK Acts that are kept anyway (their missing sections are simply absent)
KEEP_NEEDS_CHECK = [
    r"^The Code of Civil Procedure, 1908", r"^The Indian Succession Act, 1925", r"^The Legal Services Authorities Act",
    r"^The Mediation Act", r"^The Passports Act", r"^The Juvenile Justice", r"^The Transgender Persons",
    r"^The Food Safety and Standards Act", r"^The National Food Security Act",
    r"^The Employees Provident Funds", r"^The Street Vendors", r"^The Securitisation and Reconstruction",
    r"^The Human Immunodeficiency Virus", r"^The Clinical Establishments", r"^The Births, Deaths and Marriages Registration",
    r"^The Foreign Marriage Act", r"^The Police Act, 1861", r"^The KARNATAKA POLICE ACT", r"^The Unlawful Activities",
    r"^The KARNATAKA COURT-FEE", r"^The Air \(Prevention", r"^The KARNATAKA EDUCATION ACT", r"^The KARNATAKA CONTROL OF ORGANIZED",
    r"^The KARNATAKA PRIVATE MEDICAL", r"^The KARNATAKA LAND REFORMS", r"^The KARNATAKA TOWN AND COUNTRY",
    r"^The KARNATAKA LAND \(RESTRICTION ON TRANSFER",
]

DOMAIN = {  # candidate topic -> KB-style domain label
    "Family/marriage/succession/maintenance": "Family", "Labour/wages/employment": "Labour", "Consumer": "Consumer",
    "Contract/civil law": "Contract", "Property/rent/registration/land": "Property", "Housing (RERA)": "Property",
    "Money/cheques/banking disputes": "Banking and Finance", "Crime/police/courts/bail/evidence": "Criminal Law",
    "Cyber/IT/data": "Information Technology and Cyber Law", "Fundamental rights/citizenship": "Civil Rights",
    "Women/children/elderly/disability protection": "Women and Child Protection", "Motor vehicles/transport": "Motor Vehicles",
    "RTI/transparency": "Right to Information", "Education": "Education", "Health": "Health",
    "Environment/pollution": "Environment", "Local bodies (Karnataka)": "Local Government",
}

# ---------------------------------------------------------------- territorial scope
# Rule: an Act whose territorial scope is limited to State(s) / Union Territory(ies) other than Karnataka is removed.
# Decided ONLY from the Act's title and its extent clause (kb_v2_extent_scan.py reads it from the Act PDF), never from
# evaluation scores. All-India central Acts and all Karnataka Acts are kept.
ALL_INDIA = re.compile(r"(?i)whole of india|entire territor|throughout india|whole of the territor|all the states")
MENTIONS_KARNATAKA = re.compile(r"(?i)karnataka|mysore|coorg")
TERRITORIAL_NOUNS = re.compile(r"(?i)\bstates?\b|union territor|cantonment|territories (?:for the time being )?administered|presidency")
TITLE_TERRITORY = re.compile(
    r"(?i)\b(delhi|ajmer|punjab|chandigarh|goa|daman|diu|presidency|bengal|bombay|madras|oudh|central provinces|pondicherry|puducherry|"
    r"manipur|tripura|assam|sikkim|nagaland|mizoram|meghalaya|andhra|kerala|jammu|ladakh|lakshadweep|andaman)\b")
# Evidence read directly from the Acts where the extent clause is not in the extractable PDF text
TERRITORY_EVIDENCE = [
    (r"^The Public Gambling Act", "long title",
     "An Act to provide for the punishment of public gambling and the keeping of common gaming-houses in the United Provinces, East Punjab, Delhi and the Central Provinces"),
    (r"^The Presidency Small Cause Courts Act", "section 5",
     "There shall be in each of the towns of Calcutta, Madras and Bombay a Court, to be called the Court of Small Causes of Calcutta, Madras or Bombay"),
    (r"^The delhi rent control act", "extent clause (section 1(2))",
     "It extends to the areas included within the limits of the New Delhi Municipal Committee and the Delhi Cantonment Board and to such urban areas within the limits of the Municipal Corporation of Delhi as are specified in the First Schedule"),
]
# The set the classifier is expected to flag (reviewed by hand against the quotes). If the classifier ever disagrees, the
# build stops instead of silently removing or keeping an Act.
EXPECTED_TERRITORIAL_REMOVALS = {
    "The Ajmer Tenancy and Land Records Act, 1950", "The Central Provinces Tenancy Act, 1898", "The delhi rent control act, 1958",
    "The Delhi Apartment Ownership Act, 1986", "The East Punjab Urban Rent Restriction Act (Extension to Chandigarh) Act, 1974",
    "The Goa, Daman and Diu (Extension of the Code of Civil Procedure and the Arbitration Act) Act, 1965",
    "The Presidency Small Cause Courts Act, 1882", "The Public Gambling Act, 1867",
    "The Cantonments (Extension of Rent Control Laws) Act, 1957", "The Slum Areas (Improvement and Clearance) Act, 1956",
    "The Clinical Establishments (Registration and Regulation) Act, 2010",
}
# Decided by hand, overriding the extent-clause rule: the Act applies "in the first instance" to other States/UTs but also to any
# State that adopts it under Article 252, and the Act text does not say whether Karnataka did. Kept on the project owner's decision.
TERRITORY_KEEP_OVERRIDES = [r"^The Transplantation of Human Organs and Tissues Act"]


def territorial_scope(title, level, clause):
    """(limited_to_other_territory, basis, quote). Karnataka Acts are never removed."""
    if level == "Karnataka":
        return False, "", ""
    if any(re.search(p, title, flags=re.I) for p in TERRITORY_KEEP_OVERRIDES):
        return False, "", ""
    for pat, basis, quote in TERRITORY_EVIDENCE:
        if re.search(pat, title, flags=re.I):
            return True, basis, quote
    clause = clause or ""
    if clause:
        if ALL_INDIA.search(clause) or MENTIONS_KARNATAKA.search(clause):
            return False, "", ""
        if TERRITORIAL_NOUNS.search(clause):
            return True, "extent clause", clause
        return False, "", ""
    m = TITLE_TERRITORY.search(title)
    if m:
        return True, f"title ('{m.group(0)}'); extent clause not extractable", title
    return False, "", ""


# ---------------------------------------------------------------- text cleaning
SOR = re.compile(r"(?i)statement\s+of\s+objects\s+and\s+reasons")
TOC = re.compile(r"(?im)^[#*_\s>|-]*ARRANGEMENT OF SECTIONS")
ORD = r"(?:(?:FIRST|SECOND|THIRD|FOURTH|FIFTH|SIXTH|SEVENTH|EIGHTH|NINTH|TENTH|ELEVENTH|TWELFTH)\s+)?"
SCHEDULE_LINE = re.compile(r"(?im)^[#*_\s>|-]*(?:THE\s+)?" + ORD + r"SCHEDULES?\b[^\n]{0,60}$")
# CPC-style "ORDER XLII" headings: the Orders' rules are numbered like sections and trail the last section before them
ORDER_LINE = re.compile(r"(?m)^[#*_\s>|-]*ORDER\s+[IVXLCDM]+\b[^\n]{0,80}$")
FOOT_VERBS = (r"(?:Ins|Subs|Added|Rep|Omitted|Renumbered|Substituted|Inserted|The (?:words?|figures?|brackets?|letters?|marginal)|"
              r"Clauses?|Sub-sections?|Sections?|Paragraphs?|Provisos?|Explanations?|Entry|Entries|Items?|Words|Original|Now|See|Cl\.)")
FOOT_LINE = re.compile(
    r"(?m)^[>\s]*\d{1,3}\s*\.?\s*" + FOOT_VERBS +
    r"[^\n]*?(?:\bby\b|w\.e\.f|Gazette|Ordinance|\bvide\b|omitted|inserted|substituted|\bAct \d+|\bs\.\s*\d)[^\n]*$")
FOOT_SPLIT = re.compile(r"(?m)^\s*\d{1,3}\s*\.?\s*\n\s*(?:Ins|Subs)\. by[^\n]*$")
# "CHAPTER VIII" line plus the all-caps chapter title line(s) right after it
CHAPTER_HEAD = re.compile(
    r"(?m)^[\s>#*_]*(?:CHAPTER|PART)\s+[IVXLCDM\d]+[A-Z]?\b[^\n]*\n?(?:[\s>#*_]*(?=(?:[^A-Z\n]*[A-Z]){3})[^a-z\n]{3,100}\n)*")
BILL_NOTE_HEAD = re.compile(r"(?i)Bill,\s*20\d\d|inter alia,\s*provides|Statement of Objects")
START_OMITTED = re.compile(r"^[\s.\-—–\[(*_]*(omitted|repealed|rep\.\s*by|deleted)\b", re.I)
REPEAL_STUB = re.compile(r"(?i)\b(?:rep\.|omitted|repealed)\s+by\b")   # short row that only records a repeal/omission
TITLE_END = re.compile(r"\*\*|\.�|\.—|\.\s*-\s|\.\s*–")


def clean_markup(t):
    t = html.unescape(t).replace("\r", "")
    t = FOOT_SPLIT.sub("", t)
    t = FOOT_LINE.sub("", t)
    t = CHAPTER_HEAD.sub("", t)                                          # chapter headings trail the previous section
    t = re.sub(r"(?im)<br\s*/?>", " ", t)
    t = re.sub(r"(?m)^\s*\|?[\s|:\-]*\|[\s|:\-]*$", "", t)            # table separator rows
    t = t.replace("|", " ")
    t = re.sub(r"(?m)^[\s>#]*", "", t)                                  # blockquote / heading markers
    t = re.sub(r"(?m)^-\s+", "", t)                                     # bullets
    t = t.replace("**", "").replace("__", "")
    t = re.sub(r"(?<![A-Za-z0-9])_([^_\n]{1,60}?)_(?![A-Za-z0-9])", r"\1", t)   # italics: ( _a_ ) -> ( a )
    t = re.sub(r"\(\s+([^()\s]{1,6})\s+\)", r"(\1)", t)
    t = re.sub(r"\d+\s*\[", "", t)                                      # amendment footnote openers "1["
    t = t.replace("[", "").replace("]", "")
    t = re.sub(r"_{2,}", " ", t)
    t = re.sub(r"(?<![A-Za-z0-9])_|_(?![A-Za-z0-9])", "", t)          # leftover italic markers
    t = t.replace("�", "-")
    t = re.sub(r"[ \t ]+", " ", t)
    t = re.sub(r"\s*\n\s*", " ", t)
    return re.sub(r"\s{2,}", " ", t).strip()


def act_title(title):
    t = html.unescape(str(title)).replace("�", "'")
    t = re.sub(r"\s+", " ", t).strip().rstrip(".").strip()
    t = re.sub(r"^The\s+", "", t, flags=re.I)
    if t.upper() == t or sum(c.isupper() for c in t) > 0.7 * sum(c.isalpha() for c in t):
        small = {"of", "and", "the", "for", "on", "in", "to", "a", "an", "or"}
        words = []
        for i, w in enumerate(t.lower().split(" ")):
            w2 = re.sub(r"^([\(\[]?)([a-z])", lambda m: m.group(1) + m.group(2).upper(), w)
            words.append(w if (w in small and i) else w2)
        t = " ".join(words)
    t = re.sub(r"\s+,", ",", t)
    if not re.search(r",\s*\d{4}$", t):
        t = re.sub(r"\s(\d{4})$", r", \1", t)
    return t


def prefix_for(name):
    words = [w for w in re.findall(r"[A-Za-z]+", re.sub(r",\s*\d{4}$", "", name)) if w.lower() not in {"the", "of", "and", "act", "for", "in", "to", "a"}]
    return "".join(w[0].upper() for w in words)[:6] or "ACT"


def title_from(heading_rest, text, sec):
    """Section title: from the bold heading ("**66. Computer related offences.** -...") else from the first sentence."""
    rest = html.unescape(heading_rest or "")
    if rest:
        rest = re.sub(r"^[\s*_—�\-\[(]+", "", rest)
        title = TITLE_END.split(rest)[0]
        title = re.sub(r"^\d+\s*\[", "", title)
        title = re.sub(r"[*_\[\]]+", "", title).strip(" .:-—")
        if 2 < len(title) < 250:
            return title
    m = re.match(r"^\s*" + re.escape(sec) + r"\s*\.\s*(.{3,140}?)\s*[.—\-:]", text)
    return m.group(1).strip() if m else ""


def build_row_text(raw, sec, heading_rest, stats):
    """Return (title, legal_text, drop_reason). raw = section text from sections_fixed."""
    cut = None
    m = SOR.search(raw)
    if m:
        cut = m.start()
    for pat in (TOC, SCHEDULE_LINE, ORDER_LINE):
        mm = pat.search(raw)
        if mm and (cut is None or mm.start() < cut):
            cut = mm.start()
    trimmed = raw[:cut] if cut is not None else raw
    if cut is not None:
        stats["tail_trimmed"] += 1
    body = clean_markup(trimmed)
    title = title_from(heading_rest, body, sec)
    # remove the leading "66. Title.-" from the body
    body = re.sub(r"^\s*" + re.escape(sec) + r"\s*\.\s*", "", body)
    if title and body.lower().startswith(title.lower()):
        body = body[len(title):]
    body = re.sub(r"^[\s.\-—:]+", "", body)
    if title:  # the heading is sometimes repeated after stray leading text: keep what follows the last early repeat
        rep_ = re.search(re.escape(sec) + r"\s*\.\s*" + re.escape(title) + r"[\s.\-—:]*", body[:250])
        if rep_:
            body = body[rep_.end():]
    if START_OMITTED.search(body[:60]) and len(body) < 400:
        return title, "", "omitted_repealed"
    if cut is not None and len(body) < 40 and SOR.search(raw[:cut + 60]) and cut < 250:
        return title, "", "statement_of_objects_bill_note"
    if SCHEDULE_LINE.match(raw.lstrip()[:120]) and len(body) < 200:
        return title, "", "schedule"
    if BILL_NOTE_HEAD.search((title + " " + body)[:300]):
        return title, body, "statement_of_objects_bill_note"
    legal = f"{title}. {body}".strip() if title else body
    return title, legal, None


CPC = "Code of Civil Procedure, 1908"
# oversized sections that are real, important sections: split into parts at sub-section boundaries instead of dropping
SPLIT_OVERSIZED = {("Recovery Of Debts And Bankruptcy Act, 1993", "19")}
PART_CHARS = 7000


def split_parts(title, legal):
    """Split a long section at sub-section boundaries "(12) ..." into parts of about PART_CHARS characters."""
    body = legal[len(title) + 2:] if title and legal.startswith(title + ". ") else legal
    pieces = re.split(r"(?<=[.;:—-])\s+(?=\(\d+[A-Z]?\)\s)", body)
    parts, cur = [], ""
    for pc in pieces:
        if cur and len(cur) + len(pc) > PART_CHARS:
            parts.append(cur)
            cur = ""
        cur = (cur + " " + pc).strip()
    if cur:
        parts.append(cur)
    n = len(parts)
    out = []
    for i, t in enumerate(parts, 1):
        head = title if i == 1 else f"{title} (continued)"
        out.append((f"{title} (part {i} of {n})" if n > 1 else title, f"{head}. {t}"))
    return out


def load_cpc_check():
    """{(act, section): match %} from scripts/kb_v2_spot_check.py --full-act "Code of Civil Procedure, 1908"."""
    f = OIL / "kb_v2_full_act_check.csv"
    assert f.exists(), 'run: python scripts/kb_v2_spot_check.py --full-act "Code of Civil Procedure, 1908"'
    d = pd.read_csv(f)
    d = d[d.act == CPC]
    return {(r.act, str(r.section)): r.match_pct for r in d.itertuples()}


def main():
    cpc_check = load_cpc_check()
    # ---------------------------------------------------------------- own rows, verbatim
    wb = openpyxl.load_workbook(OLD_KB, read_only=True)
    ws = wb.active
    old_rows = [list(r) for r in ws.values]
    header, own = old_rows[0], old_rows[1:]
    wb.close()
    own_names = {r[header.index("act_name")] for r in own}
    print(f"own rows: {len(own)} rows, {len(own_names)} Acts")

    # ---------------------------------------------------------------- pick Acts
    rep = pd.read_csv(OIL / "act_quality_report.csv")
    cand = pd.read_csv(OIL / "candidate_acts.csv")[["act_id", "topic", "in_kb"]]
    rep = rep.merge(cand, on="act_id")
    assert (rep.in_kb == "no").all()
    keep = rep[rep.status == "OK"].copy()
    nc = rep[rep.status != "OK"]
    picked = []
    for pat in KEEP_NEEDS_CHECK:
        hit = nc[nc.act.str.contains(pat, regex=True)]
        assert len(hit) == 1, (pat, hit.act.tolist())
        picked.append(hit.iloc[0].act_id)
    keep = pd.concat([keep, nc[nc.act_id.isin(picked)]])
    print(f"open-india-law Acts before the territorial rule: {(keep.status == 'OK').sum()} OK + {(keep.status != 'OK').sum()} NEEDS CHECK = {len(keep)}")

    # ---- territorial rule: remove Acts limited to another State / Union Territory (title + extent clause only)
    scan_file = OIL / "kb_v2_extent_scan.csv"
    assert scan_file.exists(), "run: python scripts/kb_v2_extent_scan.py"
    clauses = {r.act_id: (r.extent_clause if isinstance(r.extent_clause, str) else "") for r in pd.read_csv(scan_file).itertuples()}
    src_counts = pd.read_parquet(OIL / "sections_fixed.parquet", columns=["act_id"]).act_id.value_counts()
    removed = []
    for a in keep.itertuples():
        limited, basis, quote = territorial_scope(a.act, a.level, clauses.get(a.act_id, ""))
        if limited:
            removed.append({"act": act_title(a.act), "raw_title": a.act, "act_id": a.act_id, "level": a.level, "status": a.status,
                            "sections_in_source": int(src_counts.get(a.act_id, 0)), "basis": basis, "quote": quote[:400]})
    flagged = {r["raw_title"] for r in removed}
    assert flagged == EXPECTED_TERRITORIAL_REMOVALS, (
        f"territorial classifier disagrees with the reviewed list: unexpected {sorted(flagged - EXPECTED_TERRITORIAL_REMOVALS)}, "
        f"missing {sorted(EXPECTED_TERRITORIAL_REMOVALS - flagged)}")
    keep = keep[~keep.act_id.isin({r["act_id"] for r in removed})]
    pd.DataFrame(removed).drop(columns=["raw_title"]).to_csv(OIL / "kb_v2_removed_acts.csv", index=False, encoding="utf-8-sig")
    print(f"territorial rule removed {len(removed)} Acts ({sum(r['sections_in_source'] for r in removed)} source sections); "
          f"open-india-law Acts kept: {(keep.status == 'OK').sum()} OK + {(keep.status != 'OK').sum()} NEEDS CHECK = {len(keep)}")

    urls, numbers = {}, {}
    for f in ("in_central_legislation.parquet", "in_karnataka_legislation.parquet"):
        d = pd.read_parquet(OIL / f, columns=["act_id", "source_url", "text", "chunk_id"])
        d = d.sort_values("chunk_id").drop_duplicates("act_id")
        for r in d.itertuples():
            urls[r.act_id] = r.source_url
            m = re.search(r"\(Act (\d+) of (\d{4})\)", r.text.split("\n", 1)[0])
            numbers[r.act_id] = f"{m.group(1)} of {m.group(2)}" if m else None

    secs = pd.read_parquet(OIL / "sections_fixed.parquet")
    secs = secs[secs.act_id.isin(keep.act_id)]

    # ---------------------------------------------------------------- clean
    stats = collections.Counter()
    dropped = []
    new_rows = {}
    act_meta = {}
    names_used = {}
    for a in keep.itertuples():
        name = act_title(a.act)
        assert name not in own_names, f"{name} already in own KB"
        assert name not in names_used, f"duplicate act name {name}"
        names_used[name] = a.act_id
        act_meta[a.act_id] = {"name": name, "prefix": prefix_for(name), "level": a.level, "topic": a.topic,
                              "status": a.status, "domain": DOMAIN[a.topic]}
    # law_id prefixes must be unique per Act
    by_prefix = collections.defaultdict(list)
    for k, m in act_meta.items():
        by_prefix[m["prefix"]].append(k)
    for p, ids in by_prefix.items():
        if len(ids) > 1:
            for k in ids:
                act_meta[k]["prefix"] = f"{p}{re.search(r'(\d{4})$', act_meta[k]['name']).group(1)[2:]}"
    seen_prefix = collections.Counter(m["prefix"] for m in act_meta.values())
    for k, m in act_meta.items():
        if seen_prefix[m["prefix"]] > 1:
            m["prefix"] = f"{m['prefix']}{k.split('_')[-1]}"

    for r in secs.itertuples():
        m = act_meta[r.act_id]
        stats["input_sections"] += 1
        sec = r.section_number.strip()

        def drop(rule):
            stats["dropped_" + rule] += 1
            dropped.append({"act": m["name"], "section": sec, "rule": rule, "chars": len(r.text), "excerpt": re.sub(r"\s+", " ", r.text)[:140]})

        if not re.match(r"^[1-9]\d*[A-Z]{0,3}$", sec):
            drop("bad_section_number")
            continue
        cm = clean_markup(r.text)
        # the omitted flag is per section number, so a same-numbered Schedule/rule heading can set it on a real
        # section: only trust it (and a leading "Omitted/Repealed") when the text is stub-sized
        if (r.omitted or START_OMITTED.search(cm[:120])) and len(cm) < 400:
            drop("omitted_repealed")
            continue
        if r.in_schedule and len(r.text) < 40:
            drop("schedule")
            continue
        title, legal, why = build_row_text(r.text, sec, r.heading_rest, stats)
        if why:
            drop(why)
            continue
        if len(legal) < 400 and (START_OMITTED.search(legal[:120]) or REPEAL_STUB.search(legal[len(title):len(title) + 120])):
            drop("omitted_repealed")
            continue
        if len(legal) < 40:
            drop("too_short_under_40_chars")
            continue
        if len(legal) > 20000:
            if (m["name"], sec) in SPLIT_OVERSIZED:
                stats["oversized_split_into_parts"] += 1
                stats["rows_added_by_splitting"] += len(split_parts(title, legal)) - 1
                for i, (ptitle, ptext) in enumerate(split_parts(title, legal)):
                    new_rows[(r.act_id, sec, i)] = {"act_id": r.act_id, "section": sec, "title": ptitle, "legal": ptext, "part": i}
                continue
            drop("oversized_over_20000_chars")
            continue
        if m["name"] == CPC:
            pct = cpc_check.get((CPC, sec), "missing")
            assert pct != "missing", f"CPC section {sec} was not in the full-Act check: rerun kb_v2_spot_check.py --full-act"
            if not (pct >= 90):
                drop("cpc_failed_pdf_check")
                continue
        new_rows[(r.act_id, sec, 0)] = {"act_id": r.act_id, "section": sec, "title": title, "legal": legal, "part": 0}

    # dedupe: the same text under several section numbers of one Act is a parser artefact (e.g. a table of
    # contents repeated under 25 labels), so every copy is dropped rather than keeping one wrongly-labelled copy
    def skey(sec):
        mm = re.match(r"(\d+)([A-Z]*)", sec)
        return int(mm.group(1)), mm.group(2)

    groups = collections.defaultdict(list)
    for (aid, sec, part), v in new_rows.items():
        groups[(aid, re.sub(r"\W+", "", v["legal"].lower()))].append((aid, sec, part))
    final = []
    for (aid, sec, part), v in sorted(new_rows.items(), key=lambda kv: (kv[0][0], skey(kv[0][1]), kv[0][2])):
        if len(groups[(aid, re.sub(r"\W+", "", v["legal"].lower()))]) > 1:
            stats["dropped_duplicate_text"] += 1
            dropped.append({"act": act_meta[aid]["name"], "section": sec, "rule": "duplicate_text", "chars": len(v["legal"]), "excerpt": v["legal"][:140]})
            continue
        final.append(v)

    # ---------------------------------------------------------------- write workbook
    out_header = header + ["source"]
    out = openpyxl.Workbook()
    ows = out.active
    ows.title = "Legal Knowledge Base"
    ows.append(out_header)
    for r in own:
        ows.append(r + ["own"])
    ids = set(r[0] for r in own)
    order = {aid: i for i, aid in enumerate(keep.sort_values(["level", "topic", "act"]).act_id)}
    final.sort(key=lambda v: (order[v["act_id"]], skey(v["section"]), v["part"]))
    for v in final:
        m = act_meta[v["act_id"]]
        law_id = f"{m['prefix']}-{v['section']}" + (f"-p{v['part'] + 1}" if v["part"] else "")
        assert law_id not in ids, law_id
        ids.add(law_id)
        sec_val = int(v["section"]) if v["section"].isdigit() else v["section"]
        ows.append([law_id, m["domain"], m["name"], numbers.get(v["act_id"]), sec_val, v["title"] or f"Section {v['section']}",
                    v["legal"], "Karnataka" if m["level"] == "Karnataka" else "India", "In force (per source)", None,
                    urls.get(v["act_id"]), TODAY, "Auto-imported; verify", SOURCE_OIL])
    out.save(NEW_KB)

    # ---------------------------------------------------------------- reports
    acts_df = pd.DataFrame([{**act_meta[a], "act_id": a, "sections_in_kb": sum(1 for v in final if v["act_id"] == a),
                             "source_sections": int((secs.act_id == a).sum())} for a in act_meta]).sort_values(["level", "topic", "name"])
    acts_df.to_csv(OIL / "kb_v2_acts.csv", index=False, encoding="utf-8-sig")
    pd.DataFrame(dropped).to_csv(OIL / "kb_v2_dropped_sections.csv", index=False, encoding="utf-8-sig")
    dropped_by_rule = {k[len("dropped_"):]: v for k, v in stats.items() if k.startswith("dropped_")}
    # every number below is checked against the workbook that was just written, not against in-memory counters
    check = openpyxl.load_workbook(NEW_KB, read_only=True).active
    file_rows = list(check.values)
    fh, frows = file_rows[0], file_rows[1:]
    i_act, i_src = fh.index("act_name"), fh.index("source")
    file_acts = {r[i_act] for r in frows}
    file_own = [r for r in frows if r[i_src] == "own"]
    file_new = [r for r in frows if r[i_src] != "own"]
    assert len(file_own) == len(own) and len({r[i_act] for r in file_own}) == len(own_names)
    assert len(file_new) == len(final) and len({r[i_act] for r in file_new}) == len(keep)
    assert len(frows) == len(own) + len(final) and len(file_acts) == len(own_names) + len(keep)
    assert len({r[0] for r in frows}) == len(frows), "law_id not unique in the written file"
    # accounting identity: input sections of the kept Acts = rows kept - rows created by splitting + everything dropped
    assert stats["input_sections"] == len(final) - stats["rows_added_by_splitting"] + sum(dropped_by_rule.values()), "row accounting does not add up"
    summary = {
        "own_acts": len(own_names), "own_rows": len(own),
        "open_india_law_acts_before_territorial_rule": len(keep) + len(removed),
        "acts_removed_other_state_or_ut": {
            "acts": len(removed), "source_sections": sum(r["sections_in_source"] for r in removed),
            "list": [{"act": r["act"], "basis": r["basis"], "quote": r["quote"], "source_sections": r["sections_in_source"]} for r in removed]},
        "oil_acts": len(keep), "oil_acts_ok": int((keep.status == "OK").sum()), "oil_acts_needs_check_kept": int((keep.status != "OK").sum()),
        "oil_input_sections_of_kept_acts": stats["input_sections"],
        "dropped_by_rule": dropped_by_rule, "dropped_total": sum(dropped_by_rule.values()),
        "sections_split_into_parts": stats["oversized_split_into_parts"], "rows_added_by_splitting": stats["rows_added_by_splitting"],
        "sections_with_tail_trimmed (SOR / TOC / Schedule / Order)": stats["tail_trimmed"],
        "oil_rows_kept": len(final), "total_acts": len(own_names) + len(keep), "total_rows": len(own) + len(final),
        "verified_against_written_file": {"rows": len(frows), "acts": len(file_acts), "own_rows": len(file_own), "new_rows": len(file_new),
                                          "law_ids_unique": True, "accounting_identity_holds": True},
        "built_on": TODAY,
    }
    (OIL / "kb_v2_build_report.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")
    print(json.dumps({k: v for k, v in summary.items() if k != "acts_removed_other_state_or_ut"}, indent=2))
    print(f"removed Acts: {len(removed)}")


if __name__ == "__main__":
    main()
